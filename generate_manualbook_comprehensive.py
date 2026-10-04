import os
import sys
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(13)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(11)
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
    return h

def add_p(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(11)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r_text = p.add_run(text)
    r_text.font.name = 'Arial'
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(11)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r_text = p.add_run(text)
    r_text.font.name = 'Arial'
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_img_placeholder(doc, desc):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(f"[Foto Tangkapan Layar: {desc}]")
    r.font.name = 'Arial'
    r.font.size = Pt(10)
    r.font.italic = True
    r.bold = True
    r.font.color.rgb = RGBColor(71, 85, 105)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def generate_manualbook_comprehensive():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    p_cover = doc.add_paragraph()
    p_cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cover.paragraph_format.space_after = Pt(2)
    r = p_cover.add_run("MANUAL BOOK DAN PANDUAN OPERASIONAL LENGKAP APLIKASI KUESIFY")
    r.font.name = 'Arial'
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("Panduan Penggunaan Berdasarkan Seluruh Route dan Fitur Empat Role Pengguna\nNama Tim: Kai Cenat Mewing Department (INSYFEST 2026)")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(12)
    r_sub.bold = True
    r_sub.font.color.rgb = RGBColor(51, 65, 85)

    add_heading(doc, "1. Kredensial Pengujian dan Akun Demo Juri", level=1)
    add_p(doc, "Untuk memudahkan dewan juri dalam menguji seluruh fitur platform Kuesify, berikut adalah daftar kredensial akun yang disiapkan untuk masing masing peran:")

    table = doc.add_table(rows=5, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Role Pengguna", "Email Credentials", "Password", "Fungsi dan Hak Akses Utama"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_background(cell, "1E293B")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = 'Arial'
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    creds = [
        ["Super Admin", "superadmin@kuesify.com", "password", "Pengaturan tenant global, moderasi kuis publik, dan pemantauan AI."],
        ["Org Admin", "admin@kuesify.com", "password", "Pengelolaan anggota, grup kelas, dan laporan hasil belajar."],
        ["Creator (Guru)", "creator@kuesify.com", "password", "Pembuatan kuis AI dan manual, Bank Soal, Host Live, dan koreksi essay."],
        ["Participant (Siswa)", "participant@kuesify.com", "password", "Masuk sesi kuis live via PIN, kuis mandiri, klaim badge, dan notifikasi."]
    ]
    for r_idx, row in enumerate(creds, start=1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            if c_idx == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_heading(doc, "2. Instalasi dan Pengaturan Lingkungan Pengujian Lokal", level=1)
    add_bullet(doc, "Eksekusi perintah composer install dan npm install pada direktori utama aplikasi.", bold_prefix="Langkah 1 (Dependencies): ")
    add_bullet(doc, "Pastikan file .env mengonfigurasi BROADCAST_CONNECTION=reverb dan QUEUE_CONNECTION=sync.", bold_prefix="Langkah 2 (Environment): ")
    add_bullet(doc, "Jalankan php artisan migrate:fresh opsional seed untuk membentuk struktur database awal.", bold_prefix="Langkah 3 (Database Seeding): ")
    add_bullet(doc, "Jalankan php artisan reverb:start pada terminal terpisah untuk mengaktifkan server WebSocket.", bold_prefix="Langkah 4 (WebSocket Server): ")
    add_bullet(doc, "Buka terminal baru lalu jalankan php artisan serve dan npm run dev kemudian akses http://localhost:8000.", bold_prefix="Langkah 5 (Aplikasi Server): ")

    add_heading(doc, "3. Panduan Operasional Lengkap Berdasarkan Peran Pengguna", level=1)

    # ==================== ROLE 1: PARTICIPANT ====================
    add_heading(doc, "3.1 Panduan Lengkap Role Participant (Siswa)", level=2)
    
    add_heading(doc, "A. Halaman Utama dan Gamifikasi (/dashboard)", level=3)
    add_p(doc, "1. Pengguna dapat melihat statistik XP, level pencapaian, serta grafik Daily Streak Heatmap yang mencatat keaktifan belajar harian.")
    add_p(doc, "2. Pada bilah navigasi utama, pengguna dapat mengklik profil untuk mengganti avatar atau beralih ke notifikasi.")
    add_img_placeholder(doc, "Tampilan Dashboard Utama Participant (XP Progress Bar, Daily Streak Heatmap, dan Ringkasan Aktivitas)")

    add_heading(doc, "B. Bergabung Sesi Live Quiz Multiplayer (/join)", level=3)
    add_p(doc, "1. Siswa dapat mengakses halaman /join secara langsung dari navigasi atau memasukkan 6 angka PIN kuis live dari pengajar.")
    add_p(doc, "2. Apabila siswa belum login, sistem akan meminta input nama alias dan pilihan avatar karakter dari 13 pilihan yang tersedia. Bagi siswa yang sudah terautentikasi, nama dan avatar diisi otomatis.")
    add_img_placeholder(doc, "Tampilan Halaman Join Live Quiz (Input PIN 6 Angka dan Pemilihan Avatar Character)")

    add_heading(doc, "C. Antarmuka Permainan Live Quiz Real Time (/live/sessions/{session}/play)", level=3)
    add_p(doc, "1. Peserta menyaksikan pertanyaan dan pilihan jawaban yang muncul secara serentak via Laravel Reverb WebSockets.")
    add_p(doc, "2. Setelah memilih jawaban, pilihan akan terkunci secara instan dan peserta menerima hasil berupa tampilan hijau jika Benar (dengan tambahan poin XP) atau merah jika Salah.")
    add_img_placeholder(doc, "Tampilan Layar Permainan Peserta Live Quiz (Penguncian Jawaban dan Feedback Langsung)")

    add_heading(doc, "D. Pengerjaan Kuis Mandiri dan Mode Mode Belajar (/participant/quizzes dan /attempts/{attempt}/play)", level=3)
    add_p(doc, "1. Pada halaman /participant/quizzes, siswa dapat memilih kuis mandiri atau melihat kuis terlampir dari sekolah.")
    add_p(doc, "2. Pada halaman /attempts/{attempt}/play, siswa mengerjakan kuis mandiri dengan tampilan bersih tanpa gangguan mode gelap untuk menjaga kenyamanan mata.")
    add_img_placeholder(doc, "Tampilan Layar Pengerjaan Kuis Mandiri Participant (Tampilan Bersih Light Mode)")

    add_heading(doc, "E. Eksplorasi Materi dan Catatan Mandiri (/participant/materials)", level=3)
    add_p(doc, "1. Siswa dapat membaca materi pelajaran format PDF atau PPTX pada halaman /participant/materials.")
    add_p(doc, "2. Fitur catatan interaktif memungkinkan siswa menyimpan ringkasan materi, menandai materi selesai dibaca, atau mencetak dokumen melalui tombol cetak.")
    add_img_placeholder(doc, "Tampilan Halaman Pembaca Materi Pelajaran dan Modul Catatan Mandiri Siswa")

    add_heading(doc, "F. Katalog Badge dan Notifikasi System (/participant/badges dan /notifications)", level=3)
    add_p(doc, "1. Halaman /participant/badges menampilkan koleksi lencana apresiasi yang berhasil diraih berdasarkan pencapaian skor dan aktivitas kuis.")
    add_p(doc, "2. Halaman /notifications menampilkan pesan pesan penting seperti hasil penilaian essay oleh guru, dilengkapi dengan indikator red badge pada sidebar.")
    add_img_placeholder(doc, "Tampilan Katalog Badge Prestasi dan Halaman Notifikasi Pengguna")

    # ==================== ROLE 2: CREATOR ====================
    add_heading(doc, "3.2 Panduan Lengkap Role Creator (Guru / Pengajar)", level=2)

    add_heading(doc, "A. Manajemen Kuis dan Bank Soal Adaptif (/quizzes dan /creator/question/bank)", level=3)
    add_p(doc, "1. Guru mengelola daftar kuis pada /quizzes, melakukan publikasi, pengarsipan, pembuatan salinan kuis, atau menambah kolaborator pengajar.")
    add_p(doc, "2. Pada /creator/question/bank, guru dapat membuat soal baru dengan formulir adaptif yang menyesuaikan pilihan tipe soal (Pilihan Ganda A hingga E, True/False, Isian Singkat, atau Essay).")
    add_img_placeholder(doc, "Tampilan Bank Soal Adaptif Creator (Formulir Tipe Soal Dinamis dan Pengaturan Opsi)")

    add_heading(doc, "B. Pembuat Kuis Otomatis Berbasis AI (/materials)", level=3)
    add_p(doc, "1. Guru mengunggah file dokumen pelajaran seperti PDF atau PPTX pada halaman /materials.")
    add_p(doc, "2. Dengan mengklik tombol Generasi AI, Google Gemini AI akan mengekstrak materi dan menghasilkan draf kuis lengkap dengan opsi jawaban dan pembahasan.")
    add_img_placeholder(doc, "Tampilan Ekstraksi Materi Pelajaran dan Modal Generator Kuis Berbasis AI")

    add_heading(doc, "C. Menjadi Host Sesi Live Quiz Multiplayer (/live/sessions dan /live/sessions/{session}/play)", level=3)
    add_p(doc, "1. Guru memulai sesi live quiz dari /live/sessions dan membagikan PIN kuis atau QR code ke ruang kelas.")
    add_p(doc, "2. Pada layar host, guru mengontrol jalannya permainan dengan indikator hitungan Soal ke X dari Y, penguncian jawaban otomatis, serta animasi papan peringkat Top 3.")
    add_img_placeholder(doc, "Tampilan Dashboard Host Live Quiz (Kontrol Sesi, Hitungan Soal, dan Podium Juara Top 3)")

    add_heading(doc, "D. Koreksi dan Penilaian Jawaban Essay (/attempts)", level=3)
    add_p(doc, "1. Guru memeriksa hasil pengerjaan kuis siswa pada halaman /attempts, termasuk jawaban essay yang memerlukan evaluasi manual.")
    add_p(doc, "2. Setelah memberikan nilai dan masukan, guru mengklik tombol Kirim Evaluasi yang akan secara otomatis mengirimkan notifikasi ke akun siswa.")
    add_img_placeholder(doc, "Tampilan Halaman Evaluasi Jawaban Essay dan Pengiriman Notifikasi Hasil Belajar")

    # ==================== ROLE 3: ORG ADMIN ====================
    add_heading(doc, "3.3 Panduan Lengkap Role Organization Admin (Pengelola Sekolah)", level=2)

    add_heading(doc, "A. Manajemen Anggota dan Grup Kelas (/admin/members dan /admin/groups)", level=3)
    add_p(doc, "1. Pengelola sekolah mengabaikan tombol pembuat kuis karena fokus pada administrasi organisasi.")
    add_p(doc, "2. Pada /admin/members, pengelola menambah akun guru atau siswa baru, mengubah peran, atau menonaktifkan akun yang tidak aktif.")
    add_p(doc, "3. Pada /admin/groups, pengelola mengelompokkan siswa ke dalam grup kelas untuk mempermudah distribusi kuis.")
    add_img_placeholder(doc, "Tampilan Halaman Manajemen Anggota Sekolah dan Pengelompokan Grup Kelas")

    add_heading(doc, "B. Laporan Analitik Hasil Belajar Responsif Mobile (/reports)", level=3)
    add_p(doc, "1. Halaman /reports menyediakan ringkasan nilai, tingkat kelulusan, dan distribusi skor siswa secara menyeluruh.")
    add_p(doc, "2. Antarmuka laporan dirancang responsif sehingga nyaman diakses baik melalui komputer desktop maupun perangkat ponsel pintar.")
    add_img_placeholder(doc, "Tampilan Dashboard Laporan Analitik Sekolah (Tampilan Responsif Mobile)")

    # ==================== ROLE 4: SUPER ADMIN ====================
    add_heading(doc, "3.4 Panduan Lengkap Role Super Admin (Pengelola Platform)", level=2)

    add_heading(doc, "A. Pengaturan Multi Tenant dan Pemantauan AI (/superadmin/tenants dan /superadmin/ai/monitoring)", level=3)
    add_p(doc, "1. Super admin mengelola daftar institusi sekolah pada /superadmin/tenants dan memantau status aktif setiap tenant.")
    add_p(doc, "2. Pada /superadmin/ai/monitoring, super admin memantau konsumsi token penggunaan Gemini AI dan riwayat pembuatan kuis berbasis dokumen.")
    add_img_placeholder(doc, "Tampilan Panel Kontrol Super Admin (Manajemen Tenant Sekolah dan Pemantauan AI)")

    add_heading(doc, "B. Moderasi Kuis Publik Global (/admin/moderation)", level=3)
    add_p(doc, "1. Halaman /admin/moderation menampilkan kuis kuis yang diajukan oleh pengajar untuk dipublikasikan secara umum.")
    add_p(doc, "2. Super admin memeriksa kualitas soal lalu menentukan persetujuan atau penolakan kuis publik.")
    add_img_placeholder(doc, "Tampilan Halaman Moderasi Kuis Publik Global oleh Super Admin")

    doc.save("[ManualBook] INSYFEST2026-KaiCenatMewingDepartment.docx")
    print("Comprehensive Manual Book DOCX generated successfully!")

if __name__ == "__main__":
    generate_manualbook_comprehensive()
